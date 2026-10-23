"""Read the supplied coding workbook and export reconciled website records."""
from pathlib import Path
import hashlib
import json
from collections import Counter
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / 'outputs' / 'MMR媒体报道103篇逐篇核验与阶段统计_20261014.xlsx'
workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
periods = ['1998—2000','2001—2003','2004—2009','2010—2011','2012—2014','2015—2019','2020—2026']
stances = {'支持方':'support','中立方':'neutral','误导方':'opposition','待核验':'pending'}
tags = {
    '明确说明证据不足':'insufficient',
    '引用研究或专家意见否定关联':'research',
    '报道伦理、利益或监管调查':'investigation',
    '提及MMR–自闭症可能关联':'possible',
    '将可能关联写成因果关系':'causal',
}
groups = [{'period': period, 'articles': []} for period in periods]
for row_number, row in enumerate(workbook['逐篇核验'].iter_rows(min_row=5, values_only=True), 5):
    identifier, date, period, outlet, title, url, stance, indicators, excerpt, reason, status = row
    assert stance in stances, (identifier, stance)
    coded = stances[stance]
    if coded != 'pending':
        assert excerpt, f'Classified article {identifier} has no evidence excerpt'
    labels = [] if not indicators else [tags[value] for value in indicators.split('；')]
    record = {'id': identifier, 'date': date.strftime('%Y-%m-%d'), 'outlet': outlet,
              'title': title, 'url': url, 'stance': coded, 'indicators': labels,
              'reason': reason, 'sourceRow': row_number}
    groups[periods.index(period.replace('–','—'))]['articles'].append(record)

for group in groups:
    primary = Counter(article['stance'] for article in group['articles'])
    secondary = Counter(tag for article in group['articles'] for tag in article['indicators'])
    group['counts'] = {'total':len(group['articles']), **{key:primary[key] for key in stances.values()},
                       **{key:secondary[key] for key in tags.values()}}

# Reconcile row-level calculations with the independent cached summary cells.
summary_rows = {'total':6,'pending':8,'support':10,'insufficient':12,'research':14,
                'neutral':16,'investigation':18,'possible':20,'opposition':22,'causal':24}
summary = workbook['分类统计']
for column, group in enumerate(groups, 2):
    for key, row_number in summary_rows.items():
        assert group['counts'][key] == summary.cell(row_number,column).value, (group['period'],key)
records = [article for group in groups for article in group['articles']]
assert len(records) == len({article['id'] for article in records}) == 103

metadata = {'label':'本地逐篇核验工作簿（2026-10-14）', 'workbook':SOURCE.name,
            'sheet':'逐篇核验', 'sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
            'primaryDenominator':'该阶段全部样本，含待核验',
            'secondaryDenominator':'所属立场的已分类文章数；标签可重叠',
            'oppositionMapping':'原稿“反对”位置对应核验表“误导方”，特指将可能关联写成因果关系'}
output = '// Generated from the supplied workbook by scripts/build_media_data.py.\n'
output += 'export const mediaSource = ' + json.dumps(metadata,ensure_ascii=False,indent=2) + ';\n'
output += 'export const mediaStages = ' + json.dumps(groups,ensure_ascii=False,indent=2) + ';\n'
(ROOT/'dist'/'media-data.js').write_text(output,encoding='utf-8')
print(json.dumps({'articles':len(records),'counts':[group['counts'] for group in groups]},ensure_ascii=True))
